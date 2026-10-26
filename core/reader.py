import glob
import os
from typing import List, Optional, Tuple
import openpyxl
import pandas as pd
from core.cleaner import extract_date_from_excel


class ReportChangedError(Exception):
    """Exception raised when columns in a product_list file do not match the raw_file template."""
    pass


def get_template_columns(working_file: str) -> List[str]:
    """
    Extract expected product columns from the 'raw data' sheet in working_file template.
    Returns column names excluding the initial date/day column.
    """
    wb = openpyxl.load_workbook(working_file, read_only=True, data_only=True)
    if "raw data" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet 'raw data' not found in template {working_file}")

    ws = wb["raw data"]
    headers = [str(cell.value).strip() for cell in next(ws.iter_rows(min_row=1, max_row=1)) if cell.value is not None]
    wb.close()

    if not headers:
        raise ValueError(f"No headers found in 'raw data' sheet of {working_file}")

    # Column 0 is 'day'/'date' in raw data sheet, remaining columns are product metrics
    if headers[0].lower() in ("day", "date"):
        return headers[1:]
    return headers


def validate_product_columns(
    product_cols: List[str],
    expected_cols: List[str],
    file_name: str = "",
    template_name: str = "",
) -> None:
    """
    Validate that product_list file columns match the raw_file template columns.
    Raises ReportChangedError if there is any mismatch.
    """
    cleaned_product = [str(c).strip() for c in product_cols]
    cleaned_expected = [str(c).strip() for c in expected_cols]

    if cleaned_product != cleaned_expected:
        missing = [c for c in cleaned_expected if c not in cleaned_product]
        extra = [c for c in cleaned_product if c not in cleaned_expected]

        details = []
        if missing:
            details.append(f"Missing columns: {missing}")
        if extra:
            details.append(f"Unexpected extra columns: {extra}")
        if not missing and not extra and cleaned_product != cleaned_expected:
            details.append("Column order is different")

        detail_msg = "\n  - " + "\n  - ".join(details) if details else ""
        file_info = f" in '{file_name}'" if file_name else ""
        template_info = f" template '{template_name}'" if template_name else " template"

        raise ReportChangedError(
            f"Report Changed: Columns{file_info} do not match{template_info}.{detail_msg}"
        )


def load_and_combine_product_files(
    folder_path: str,
    working_file: Optional[str] = None,
) -> Tuple[pd.DataFrame, int]:
    """
    Search for product_list*.xlsx files in folder_path, validate columns against working_file template,
    extract date, read records (header at row 3, ID column as string), and combine into a single DataFrame.
    """
    search_pattern = os.path.join(folder_path, "product_list*.xlsx")
    product_files = glob.glob(search_pattern)

    if not product_files:
        return pd.DataFrame(), 0

    expected_cols = None
    if working_file and os.path.isfile(working_file):
        expected_cols = get_template_columns(working_file)

    all_data = []

    for file in product_files:
        try:
            date_val = extract_date_from_excel(file)

            # Read with header at row index 2 (row 3 in Excel) to preserve column names
            # Force first column (Column 0 - ID) to string to prevent truncation
            df_data = pd.read_excel(file, header=2, dtype={0: str})

            # Validate that columns match the raw_file template
            if expected_cols:
                validate_product_columns(
                    list(df_data.columns),
                    expected_cols,
                    file_name=os.path.basename(file),
                    template_name=os.path.basename(working_file),
                )

            df_data.insert(0, "Day", date_val)
            all_data.append(df_data)
        except ReportChangedError:
            # Re-raise ReportChangedError immediately to halt execution
            raise
        except Exception as e:
            print(f"Error when reading {os.path.basename(file)}: {e}")

    if not all_data:
        return pd.DataFrame(), len(product_files)

    combined_df = pd.concat(all_data, ignore_index=True)

    # Sort descending by Day
    combined_df["Sort_Date"] = pd.to_datetime(combined_df["Day"], errors="coerce")
    combined_df = combined_df.sort_values(by="Sort_Date", ascending=False).drop(columns=["Sort_Date"])

    return combined_df, len(product_files)

