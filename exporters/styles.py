from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Headers styling
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Total summary row styling
TOTAL_FONT = Font(name="Calibri", size=11, bold=True)
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="595959"),
    bottom=Side(style="double", color="595959"),
)

# Common cell borders & alignments
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Number formats
FORMAT_PERCENTAGE = "0.00%"
FORMAT_NUMBER = "#,##0"

# Row heights
HEADER_ROW_HEIGHT = 28
DATA_ROW_HEIGHT = 20
TOTAL_ROW_HEIGHT = 22
