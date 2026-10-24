from typing import List
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

from core.processor import calculate_total_summary
from exporters.styles import (
    HEADER_FONT,
    HEADER_FILL,
    HEADER_ALIGNMENT,
    THIN_BORDER,
    TOTAL_FONT,
    TOTAL_FILL,
    TOTAL_BORDER,
    ALIGN_CENTER,
    ALIGN_RIGHT,
    FORMAT_PERCENTAGE,
    FORMAT_NUMBER,
    HEADER_ROW_HEIGHT,
    DATA_ROW_HEIGHT,
    TOTAL_ROW_HEIGHT,
)


def get_month_suffix(dates_series: pd.Series, default_suffix: str = "Sep26") -> str:
    """Extract month-year suffix (e.g. Sep26) from a series of date values."""
    valid_dates = pd.to_datetime(dates_series, errors="coerce").dropna()
    if not valid_dates.empty:
        return valid_dates.iloc[0].strftime("%b%y")
    return default_suffix


def export_raw_data(working_file: str, combined_df: pd.DataFrame, output_raw_file: str) -> None:
    """Populate raw data sheet into template Excel file and save to output file."""
    wb = load_workbook(working_file)

    if "raw data" not in wb.sheetnames:
        raise ValueError("Error: The 'raw data' tab was not found in the working file.")

    ws = wb["raw data"]
    start_col = 1
    start_row = 2

    for r_idx, row in enumerate(combined_df.values, start=start_row):
        for c_idx, value in enumerate(row, start=start_col):
            if pd.isna(value):
                ws.cell(row=r_idx, column=c_idx, value="")
            else:
                ws.cell(row=r_idx, column=c_idx, value=value)

    wb.save(output_raw_file)


def export_daily_summary_to_excel(
    daily_summary: pd.DataFrame, final_cols: List[str], output_summary_file: str
) -> None:
    """Export daily summary DataFrame to an Excel workbook with styling and formatted numbers."""
    wb_sum = openpyxl.Workbook()
    ws_sum = wb_sum.active
    ws_sum.title = "Daily Summary"

    # Write headers
    ws_sum.row_dimensions[1].height = HEADER_ROW_HEIGHT
    for c_idx, col_name in enumerate(final_cols, start=1):
        cell = ws_sum.cell(row=1, column=c_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Write data rows
    for r_idx, row in enumerate(daily_summary.itertuples(index=False), start=2):
        ws_sum.row_dimensions[r_idx].height = DATA_ROW_HEIGHT
        for c_idx, (col_name, val) in enumerate(zip(final_cols, row), start=1):
            cell = ws_sum.cell(row=r_idx, column=c_idx, value=val)
            cell.border = THIN_BORDER
            if col_name == "Day":
                cell.alignment = ALIGN_CENTER
            elif "rate" in col_name.lower():
                cell.number_format = FORMAT_PERCENTAGE
                cell.alignment = ALIGN_RIGHT
            else:
                cell.number_format = FORMAT_NUMBER
                cell.alignment = ALIGN_RIGHT

    # Add Total summary row at bottom
    tot_row_idx = len(daily_summary) + 2
    ws_sum.row_dimensions[tot_row_idx].height = TOTAL_ROW_HEIGHT
    tot_dict = calculate_total_summary(daily_summary, final_cols)

    for c_idx, col_name in enumerate(final_cols, start=1):
        val = tot_dict.get(col_name, "")
        cell = ws_sum.cell(row=tot_row_idx, column=c_idx, value=val)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        if col_name == "Day":
            cell.alignment = ALIGN_CENTER
        elif "rate" in col_name.lower():
            cell.number_format = FORMAT_PERCENTAGE
            cell.alignment = ALIGN_RIGHT
        else:
            cell.number_format = FORMAT_NUMBER
            cell.alignment = ALIGN_RIGHT

    # Auto-fit column widths
    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze header row and Day column
    ws_sum.freeze_panes = "B2"

    wb_sum.save(output_summary_file)
